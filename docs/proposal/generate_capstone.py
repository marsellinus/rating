# -*- coding: utf-8 -*-
"""
Generate the Capstone Project proposal (Prodi Informatika, Universitas Siliwangi)
as a properly formatted MS Word (.docx) document.

Structure follows "Pedoman Pelaksanaan Capstone Project" section E.3 (Penyusunan
Proposal): 11 required points, plus WBS/schedule/risk from the RPS.
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "docs/proposal/Capstone-Project-RATIG.docx"
UI = "docs/proposal/ui"
FONT = "Times New Roman"

doc = Document()

# ---------------------------------------------------------------- page setup
for section in doc.sections:
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(4.0)
    section.left_margin = Cm(4.0)
    section.bottom_margin = Cm(3.0)
    section.right_margin = Cm(3.0)

# ---------------------------------------------------------------- base style
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(12)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
pf = normal.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
pf.space_after = Pt(6)


def _set_run(run, size=12, bold=False, italic=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def para(text="", size=12, bold=False, italic=False, align=None,
         space_after=6, space_before=0, indent_left=None, first_line=None,
         line=1.5):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line
    if indent_left is not None:
        p.paragraph_format.left_indent = Cm(indent_left)
    if first_line is not None:
        p.paragraph_format.first_line_indent = Cm(first_line)
    if text:
        r = p.add_run(text)
        _set_run(r, size=size, bold=bold, italic=italic)
    return p


def rich(parts, align=None, space_after=6, space_before=0, indent_left=None,
         first_line=None, line=1.5):
    """parts: list of (text, bold, italic)."""
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line
    if indent_left is not None:
        p.paragraph_format.left_indent = Cm(indent_left)
    if first_line is not None:
        p.paragraph_format.first_line_indent = Cm(first_line)
    for t, b, i in parts:
        r = p.add_run(t)
        _set_run(r, bold=b, italic=i)
    return p


def h1(text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    _set_run(r, size=14, bold=True, color=RGBColor(0, 0, 0))
    return p


def h2(text):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    _set_run(r, size=12, bold=True, color=RGBColor(0, 0, 0))
    return p


def bullets(items, numbered=False):
    for it in items:
        p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.5
        # A single (text, bold, italic) run spec.
        if isinstance(it, tuple) and len(it) == 3 and isinstance(it[0], str):
            r = p.add_run(it[0])
            _set_run(r, bold=it[1], italic=it[2])
        # A list of (text, bold, italic) parts.
        elif isinstance(it, (list, tuple)):
            for t, b, i in it:
                r = p.add_run(t)
                _set_run(r, bold=b, italic=i)
        else:
            r = p.add_run(it)
            _set_run(r)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htxt in enumerate(headers):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(htxt)
        _set_run(r, size=11, bold=True)
        hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _shade(hdr[i], "D9EAD3")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            _set_run(r, size=11)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run(text)
    _set_run(r, size=10, italic=True)


def figure(path, cap, width_cm=9.5):
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(path, width=Cm(width_cm))
        caption(cap)


def figrow(items, width_cm=5.4):
    """items: list of (path, caption). Rendered side by side in a borderless row."""
    items = [(p, c) for p, c in items if os.path.exists(p)]
    if not items:
        return
    t = doc.add_table(rows=1, cols=len(items))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for cell, (path, cap) in zip(t.rows[0].cells, items):
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(2)
        cp.add_run().add_picture(path, width=Cm(width_cm))
        cp2 = cell.add_paragraph()
        cp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp2.paragraph_format.space_after = Pt(6)
        r = cp2.add_run(cap)
        _set_run(r, size=9, italic=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_toc():
    p = doc.add_paragraph()
    run = p.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r'TOC \o "1-3" \h \z \u')
    t = OxmlElement("w:t")
    t.text = "Klik kanan lalu pilih “Update Field” untuk membuat daftar isi."
    r2 = OxmlElement("w:r")
    r2.append(t)
    fld.append(r2)
    run._r.addnext(fld)


# ================================================================ COVER
para("PROPOSAL CAPSTONE PROJECT", size=12, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

para("RATIG: Sistem Pemeriksaan Kelelahan Pekerja Berbasis Waktu Reaksi "
     "pada Perangkat Android dengan Dukungan Offline-First",
     size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

para("(Reaction Time Fatigue)", size=12, italic=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30)

para("Disusun oleh:", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
para("Kelompok Capstone Project", size=12, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

tbl = doc.add_table(rows=0, cols=3)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
members = [
    ("1.", "Nama Anggota 1", "0000000001"),
    ("2.", "Nama Anggota 2", "0000000002"),
    ("3.", "Nama Anggota 3", "0000000003"),
]
for no, nm, nim in members:
    c = tbl.add_row().cells
    c[0].text = ""; c[1].text = ""; c[2].text = ""
    for cell, txt, al in ((c[0], no, WD_ALIGN_PARAGRAPH.RIGHT),
                          (c[1], nm, WD_ALIGN_PARAGRAPH.LEFT),
                          (c[2], nim, WD_ALIGN_PARAGRAPH.LEFT)):
        r = cell.paragraphs[0].add_run(txt)
        _set_run(r)
        cell.paragraphs[0].alignment = al
        cell.paragraphs[0].paragraph_format.line_spacing = 1.5
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)

para("", space_after=24)
para("PROGRAM STUDI INFORMATIKA", size=13, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("FAKULTAS TEKNIK", size=13, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("UNIVERSITAS SILIWANGI", size=13, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("2026", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_page_break()

# ================================================================ TOC
h1("DAFTAR ISI")
add_toc()
doc.add_page_break()

# ================================================================ 1. JUDUL
h1("BAB I  PENDAHULUAN")

h2("1.1  Judul Proyek")
rich([("Judul: ", True, False),
      ("RATIG: Sistem Pemeriksaan Kelelahan Pekerja Berbasis Waktu Reaksi pada "
       "Perangkat Android dengan Dukungan Offline-First", False, False)])
rich([("Kelompok Keilmuan (KK): ", True, False),
      ("Teknologi Multimedia dan Game (TMG), lintas KK dengan Intelligent "
       "Systems and Informatics (ISI) pada bagian analisis metrik dan "
       "klasifikasi hasil.", False, False)])
rich([("Bidang: ", True, False),
      ("Aplikasi mobile, interaksi pengguna, dan sistem pendukung keputusan.", False, False)])

# ================================================================ 1.2 LATAR
h2("1.2  Latar Belakang dan Identifikasi Masalah")
para("Kelelahan (fatigue) merupakan salah satu penyebab utama kecelakaan kerja "
     "di sektor industri. Kondisi kelelahan menurunkan kewaspadaan, memperlambat "
     "pengambilan keputusan, dan meningkatkan peluang terjadinya kesalahan "
     "operasional. Salah satu indikator objektif yang dapat diukur untuk "
     "mendeteksi kelelahan adalah waktu reaksi (reaction time), yaitu selang "
     "waktu antara munculnya stimulus dan respons yang diberikan oleh pekerja.",
     first_line=1.0)
para("Permasalahan yang teridentifikasi di lapangan antara lain: (1) penilaian "
     "kelelahan masih banyak dilakukan secara subjektif melalui wawancara atau "
     "kuesioner sehingga sulit dibandingkan antar waktu dan antar pekerja; "
     "(2) proses pemeriksaan manual memerlukan waktu dan tenaga, sementara jumlah "
     "pekerja yang harus diperiksa relatif banyak; (3) pencatatan hasil masih "
     "menggunakan kertas atau lembar kerja sehingga pemantauan menyeluruh dan "
     "penelusuran riwayat menjadi sulit; serta (4) lokasi kerja sering berada di "
     "area dengan koneksi internet yang tidak stabil, sehingga aplikasi berbasis "
     "jaringan tidak dapat diandalkan.", first_line=1.0)
para("Berdasarkan kondisi tersebut, diperlukan sebuah sistem yang mampu "
     "melakukan pemeriksaan kelelahan secara objektif dan cepat, dapat digunakan "
     "untuk banyak pekerja dalam satu perangkat, tetap berfungsi tanpa koneksi "
     "internet, serta menyimpan dan menampilkan hasil secara aman dan tertelusur. "
     "Solusi yang diusulkan adalah aplikasi Android bernama RATIG (Reaction Time "
     "Fatigue) dengan pendekatan offline-first.", first_line=1.0)

# ================================================================ 1.3 RUMUSAN
h2("1.3  Rumusan Masalah")
bullets([
    "Bagaimana merancang aplikasi pemeriksaan kelelahan berbasis waktu reaksi "
    "yang mengukur waktu reaksi secara akurat pada perangkat Android?",
    "Bagaimana merancang mekanisme identifikasi pekerja yang cepat dan aman "
    "untuk pola pemeriksaan banyak pekerja dalam satu perangkat?",
    "Bagaimana merancang arsitektur offline-first sehingga pemeriksaan tetap "
    "dapat dilakukan tanpa koneksi internet dan data tersinkronisasi secara "
    "andal saat kembali online?",
    "Bagaimana menyajikan hasil pemeriksaan, klasifikasi, dan pemantauan secara "
    "aman dengan pembatasan akses berdasarkan peran pengguna?",
], numbered=True)

# ================================================================ 1.4 TUJUAN
h2("1.4  Tujuan")
bullets([
    "Membangun aplikasi Android RATIG yang mampu mengukur waktu reaksi pekerja "
    "secara akurat menggunakan monotonic clock.",
    "Menyediakan mekanisme identifikasi pekerja melalui NIK dan pemindai "
    "barcode/QR untuk mendukung pemeriksaan banyak pekerja.",
    "Menerapkan arsitektur offline-first dengan sinkronisasi otomatis dan "
    "idempotent ke basis data server.",
    "Menyediakan pengelolaan pengguna berbasis peran serta dasbor pemantauan "
    "dan pelaporan hasil pemeriksaan.",
    "Mengevaluasi fungsionalitas dan usability aplikasi melalui pengujian dan "
    "evaluasi pengguna.",
], numbered=True)

# ================================================================ 1.5 MANFAAT
h2("1.5  Manfaat")
rich([("a.  Bagi Pekerja: ", True, False),
      ("memperoleh pemeriksaan kelelahan yang objektif, cepat, dan tidak "
       "mengganggu waktu kerja secara berlebihan.", False, False)])
rich([("b.  Bagi Petugas Pemeriksa (Paramedis/SHE): ", True, False),
      ("memperoleh alat bantu keputusan yang praktis, dapat digunakan di lokasi "
       "tanpa jaringan, dan mempermudah pendokumentasian hasil.", False, False)])
rich([("c.  Bagi Manajemen Perusahaan: ", True, False),
      ("memperoleh data kelelahan yang tertelusur untuk pemantauan tren per "
       "departemen, shift, dan waktu sebagai dasar tindakan pencegahan.", False, False)])
rich([("d.  Bagi Program Studi: ", True, False),
      ("menjadi produk Capstone yang menunjukkan penerapan ilmu rekayasa "
       "perangkat lunak, basis data, interaksi pengguna, dan keamanan sistem.", False, False)])

# ================================================================ 1.6 BATASAN
h2("1.6  Batasan Masalah")
bullets([
    "Aplikasi dibangun untuk platform Android (minimum Android 8.0 / API 26).",
    "Pemeriksaan difokuskan pada pengukuran waktu reaksi; hasil bersifat alat "
    "bantu keputusan, bukan diagnosis medis.",
    "Klasifikasi kelelahan mengikuti aturan/ambang yang ditetapkan pengguna "
    "sistem dan tidak menggantikan penilaian pemeriksa yang berwenang.",
    "Pemindai barcode/QR menggunakan kamera perangkat; perangkat keras khusus "
    "tidak termasuk dalam lingkup.",
], numbered=True)

# ================================================================ BAB II
doc.add_page_break()
h1("BAB II  TINJAUAN PUSTAKA")

h2("2.1  Kelelahan Kerja dan Waktu Reaksi")
para("Kelelahan kerja adalah kondisi penurunan kapasitas fisik maupun mental "
     "akibat aktivitas yang berkepanjangan. Salah satu indikator objektif yang "
     "umum digunakan adalah waktu reaksi, karena kelelahan memperlambat proses "
     "pengolahan informasi dan pengambilan keputusan. Pengukuran waktu reaksi "
     "yang baik memerlukan ketepatan waktu yang tinggi; oleh karena itu aplikasi "
     "ini menggunakan jam monotonik (monotonic clock) yang tidak terpengaruh "
     "perubahan jam sistem, alih-alih jam dinding.", first_line=1.0)

h2("2.2  Benchmark Solusi")
para("Beberapa pendekatan yang telah ada di antaranya: (a) aplikasi pengukuran "
     "waktu reaksi sederhana yang hanya menghasilkan satu jenis tes dan tidak "
     "mendukung banyak pekerja; (b) sistem pemeriksaan berbasis komputer yang "
     "bergantung pada jaringan dan perangkat tetap; serta (c) pencatatan manual "
     "menggunakan kertas atau lembar kerja. Keterbatasan umum dari solusi yang ada "
     "adalah ketiadaan dukungan offline, keterbatasan identifikasi pekerja, dan "
     "minimnya pembatasan akses serta jejak audit.", first_line=1.0)
para("RATIG membedakan diri melalui kombinasi: beberapa mode pemeriksaan dalam "
     "satu aplikasi, identifikasi cepat berbasis NIK/barcode, arsitektur "
     "offline-first dengan sinkronisasi idempotent, pengelolaan akses berbasis "
     "peran, dan hasil yang tertelusur (audit).", first_line=1.0)

h2("2.3  Landasan Teori")
bullets([
    ("Rekayasa Perangkat Lunak: ", True, False),
    ("pengembangan iteratif dan inkremental dengan pemisahan lapisan (presentasi, "
     "domain, data) untuk menjaga kualitas dan kemudahan pemeliharaan "
     "(Sommerville).", False, False),
    ("Arsitektur Offline-First: ", True, False),
    ("penyimpanan data lokal sebagai sumber kerja utama dengan sinkronisasi "
     "latar belakang yang idempotent dan tahan gangguan jaringan.", False, False),
    ("Interaksi Manusia dan Komputer (IMK): ", True, False),
    ("penerapan Material Design dan evaluasi usability untuk memastikan "
     "kemudahan penggunaan di lapangan.", False, False),
    ("Keamanan dan Kendali Akses: ", True, False),
    ("penerapan Row Level Security pada basis data dan prinsip hak akses "
     "minimum untuk melindungi data pribadi pekerja.", False, False),
    ("Pengujian Perangkat Lunak: ", True, False),
    ("pengujian unit, pengujian fungsional, dan evaluasi penerimaan pengguna "
     "(UAT) sebagai bukti terukur kualitas produk.", False, False),
], numbered=True)

# ================================================================ BAB III
doc.add_page_break()
h1("BAB III  METODOLOGI PELAKSANAAN")

h2("3.1  Metode Pengembangan")
para("Proyek dilaksanakan dengan model Project Based Learning (PjBL) dan "
     "pendekatan pengembangan iteratif/Agile (sprint). Setiap iterasi mencakup "
     "perencanaan, perancangan, implementasi, pengujian, dan evaluasi. Praktik "
     "rekayasa yang diterapkan meliputi penggunaan version control (Git), "
     "pengelolaan issue, dan dokumentasi teknis.", first_line=1.0)

h2("3.2  Tahapan Pelaksanaan")
table(
    ["Tahap", "Kegiatan", "Luaran"],
    [
        ["1. Problem Discovery", "Identifikasi masalah, pengguna, dan kebutuhan; penetapan indikator keberhasilan.", "Problem statement, kebutuhan pengguna"],
        ["2. Analisis Kebutuhan", "Penyusunan kebutuhan fungsional & nonfungsional, batasan, dan lingkup.", "Spesifikasi kebutuhan"],
        ["3. Perancangan", "Perancangan arsitektur, model data, alur proses, dan UI/UX.", "Dokumen desain, prototipe UI"],
        ["4. Implementasi", "Pembangunan aplikasi Android, basis data, dan fungsi server secara iteratif.", "Aplikasi (MVP → beta)"],
        ["5. Integrasi & Deployment", "Integrasi komponen, konfigurasi, dan quality assurance.", "Produk terintegrasi"],
        ["6. Pengujian & Evaluasi", "Pengujian fungsional, unit test, dan evaluasi usability (UAT).", "Bukti pengujian, hasil UAT"],
        ["7. Dokumentasi & Diseminasi", "Laporan teknis, manual pengguna, dan presentasi demo.", "Laporan, manual, materi demo"],
    ],
    widths=[4.0, 6.5, 4.0],
)

h2("3.3  Arsitektur Sistem")
para("Sistem terdiri atas aplikasi klien Android dan layanan backend. Aplikasi "
     "klien dibangun dengan Kotlin dan Jetpack Compose, menyimpan data secara "
     "lokal (Room) sebagai basis kerja offline-first, dan melakukan sinkronisasi "
     "latar belakang melalui WorkManager. Backend menggunakan Supabase "
     "(PostgreSQL) dengan Row Level Security, autentikasi, serta Edge Functions "
     "untuk operasi yang memerlukan hak istimewa.", first_line=1.0)
table(
    ["Komponen", "Teknologi", "Fungsi"],
    [
        ["Antarmuka", "Jetpack Compose (Material 3)", "Tampilan dan interaksi pengguna"],
        ["Penyimpanan lokal", "Room (SQLite)", "Basis kerja offline-first"],
        ["Sinkronisasi", "WorkManager", "Sinkronisasi latar belakang idempotent"],
        ["Pemindai", "CameraX + ML Kit", "Pembacaan barcode/QR NIK"],
        ["Backend", "Supabase (PostgreSQL, Auth)", "Data, autentikasi, kendali akses"],
        ["Keamanan", "Row Level Security", "Pembatasan akses per peran"],
    ],
    widths=[3.6, 5.4, 5.5],
)

h2("3.4  Rancangan Pengujian")
bullets([
    "Pengujian unit (unit test) untuk logika perhitungan metrik, klasifikasi, "
    "validasi identitas, dan penanganan kegagalan sinkronisasi.",
    "Pengujian fungsional untuk alur utama: identifikasi, pelaksanaan tes, "
    "finalisasi hasil, dan sinkronisasi.",
    "Evaluasi usability dan penerimaan pengguna (UAT) menggunakan kuesioner "
    "terhadap petugas pemeriksa.",
    "Pengujian ketahanan offline: menjalankan pemeriksaan tanpa jaringan lalu "
    "memastikan data tersinkronisasi saat online.",
], numbered=True)

# ================================================================ 3.5 WBS
h2("3.5  Work Breakdown Structure (WBS)")
table(
    ["No", "Paket Kerja", "Sub-Pekerjaan Utama"],
    [
        ["1", "Manajemen Proyek", "Perencanaan, jadwal, logbook, manajemen risiko"],
        ["2", "Analisis & Desain", "Kebutuhan, arsitektur, model data, UI/UX"],
        ["3", "Implementasi Klien", "UI Compose, mesin waktu reaksi, mode tes, identifikasi"],
        ["4", "Implementasi Offline & Sinkronisasi", "Basis data lokal, antrean sinkronisasi, status sinkron"],
        ["5", "Implementasi Backend", "Skema basis data, RLS, RPC, Edge Functions"],
        ["6", "Pengujian & Evaluasi", "Unit test, fungsional, UAT"],
        ["7", "Dokumentasi & Diseminasi", "Laporan, manual pengguna, poster/video, demo"],
    ],
    widths=[1.2, 5.0, 8.3],
)

# ================================================================ 3.6 JADWAL
h2("3.6  Rencana Kegiatan (16 Minggu)")
table(
    ["Minggu", "Kegiatan", "Luaran"],
    [
        ["1–3", "Orientasi, problem discovery, analisis kebutuhan", "Problem statement, kebutuhan"],
        ["4–5", "Benchmark, penyusunan proposal, seminar proposal", "Proposal"],
        ["6–8", "Perancangan arsitektur, data, UI/UX, design review", "Dokumen desain"],
        ["9–10", "Implementasi Sprint 1 & 2 (MVP dan fitur utama)", "Aplikasi MVP"],
        ["11", "Integrasi, deployment, quality assurance", "Beta prototype"],
        ["12–13", "Pengujian fungsional, monitoring, UAT, penyempurnaan", "Bukti uji, hasil UAT"],
        ["14", "Manajemen proyek, logbook, peer assessment", "Logbook, laporan kontribusi"],
        ["15–16", "Finalisasi, dokumentasi, presentasi/demo", "Produk final, laporan"],
    ],
    widths=[1.8, 8.2, 4.5],
)

# ================================================================ 3.7 RISIKO
h2("3.7  Risiko dan Mitigasi")
table(
    ["Risiko", "Dampak", "Mitigasi"],
    [
        ["Jadwal bertabrakan dengan mata kuliah lain", "Keterlambatan milestone", "Perencanaan sprint ringkas, pembagian tugas jelas, monitoring mingguan"],
        ["Perbedaan waktu antar perangkat memengaruhi akurasi", "Hasil tidak konsisten", "Penggunaan monotonic clock dan validasi rentang waktu"],
        ["Koneksi internet tidak stabil", "Data tidak tersinkron", "Arsitektur offline-first dan sinkronisasi ulang otomatis"],
        ["Kesalahan/kehilangan data pekerja", "Masalah privasi", "Kendali akses berbasis peran, masking NIK, jejak audit"],
        ["Ruang lingkup melebar", "Proyek tidak selesai", "Penetapan batasan dan prioritas fitur inti (MVP)"],
    ],
    widths=[4.6, 4.0, 5.9],
)

# ================================================================ 3.8 SUMBER DAYA
h2("3.8  Kebutuhan Sumber Daya")
table(
    ["Kategori", "Kebutuhan"],
    [
        ["Perangkat", "Laptop untuk pengembangan; perangkat Android untuk pengujian (kamera)"],
        ["Perangkat lunak", "Android Studio, Git, akun Supabase (lingkungan uji)"],
        ["Data", "Data uji pekerja/departemen/shift (anonim) dan aturan klasifikasi contoh"],
        ["Manusia", "3 mahasiswa anggota kelompok"],
    ],
    widths=[3.4, 11.1],
)

# ================================================================ 3.9 TAMPILAN
h2("3.9  Rancangan Antarmuka (Cuplikan)")
para("Berikut cuplikan antarmuka yang telah dirancang dengan Material Design 3.",
     first_line=1.0)
figrow([
    (f"{UI}/02_dashboard_superadmin.png", "Gambar 3.1  Dasbor pemantauan"),
    (f"{UI}/06_pengguna_semua.png", "Gambar 3.2  Pengelolaan pengguna"),
])
figrow([
    (f"{UI}/11_detail_sesi.png", "Gambar 3.3  Detail hasil pemeriksaan"),
    (f"{UI}/20_dashboard_user.png", "Gambar 3.4  Beranda pengguna"),
])

# ================================================================ BAB IV
doc.add_page_break()
h1("BAB IV  PEMBAGIAN TUGAS DAN LUARAN")

h2("4.1  Pembagian Tugas Anggota")
table(
    ["Anggota", "Peran", "Tanggung Jawab Utama"],
    [
        ["Anggota 1\n(Ketua)", "Arsitektur & Backend",
         "Perancangan arsitektur dan model data, implementasi backend (skema basis data, "
         "kebijakan akses, fungsi server), koordinasi tim dan manajemen proyek."],
        ["Anggota 2", "Android Developer & UI/UX",
         "Perancangan antarmuka (Material Design 3), implementasi aplikasi (mesin waktu "
         "reaksi, mode tes, identifikasi NIK/barcode)."],
        ["Anggota 3", "Offline/Sinkronisasi, QA & Dokumentasi",
         "Implementasi basis data lokal dan sinkronisasi, penyusunan pengujian, evaluasi "
         "usability (UAT), serta dokumentasi dan manual pengguna."],
    ],
    widths=[2.8, 4.0, 7.7],
)
para("Setiap anggota berkontribusi pada keseluruhan proyek dan bertanggung jawab "
     "atas keberhasilan proyek secara tim.", italic=True, first_line=1.0)

h2("4.2  Luaran yang Ditargetkan")
bullets([
    "Produk/prototipe aplikasi Android RATIG yang berfungsi (MVP hingga beta).",
    "Repositori kode sumber beserta riwayat perubahan (version control).",
    "Laporan teknis dan dokumentasi sistem.",
    "Manual pengguna dan manual instalasi.",
    "Bukti pengujian (unit test, pengujian fungsional) dan hasil evaluasi "
    "usability (UAT).",
    "Materi diseminasi: poster dan/atau video demonstrasi produk.",
], numbered=True)

h2("4.3  Indikator Keberhasilan")
table(
    ["Indikator", "Target"],
    [
        ["Akurasi pengukuran waktu reaksi", "Konsisten dan tahan terhadap perubahan jam sistem (monotonic clock)"],
        ["Keberhasilan sinkronisasi", "Data pemeriksaan tersinkronisasi otomatis saat kembali online"],
        ["Kelengkapan alur utama", "Identifikasi → tes → hasil → finalisasi berjalan tanpa kesalahan"],
        ["Evaluasi usability", "Skor penerimaan pengguna pada kategori baik"],
        ["Kualitas rekayasa", "Unit test lulus dan kode terkelola dalam repositori"],
    ],
    widths=[5.5, 9.0],
)

# ================================================================ BAB V
h1("BAB V  PENUTUP")
para("Proposal ini menguraikan rencana pelaksanaan Capstone Project berupa "
     "pengembangan aplikasi RATIG, yaitu sistem pemeriksaan kelelahan pekerja "
     "berbasis waktu reaksi pada perangkat Android dengan dukungan offline-first. "
     "Proyek diarahkan untuk menghasilkan produk yang berfungsi, teruji, dan "
     "bermanfaat bagi pengguna, dengan tetap memperhatikan aspek keamanan, "
     "privasi, dan etika. Melalui pelaksanaan proyek ini, anggota kelompok "
     "diharapkan dapat menerapkan kompetensi rekayasa perangkat lunak secara "
     "kolaboratif dan profesional.", first_line=1.0)
para("Hasil RATIG bersifat alat bantu keputusan dan tidak menggantikan "
     "penilaian medis maupun keputusan resmi perusahaan.", italic=True,
     first_line=1.0)

# ================================================================ REFERENSI
h1("REFERENSI")
refs = [
    "Sommerville, I. Software Engineering. Pearson.",
    "Kurose, J. F., & Ross, K. W. Computer Networking: A Top-Down Approach. Pearson.",
    "Stallings, W. Cryptography and Network Security: Principles and Practice. Pearson.",
    "Russell, S., & Norvig, P. Artificial Intelligence: A Modern Approach. Pearson.",
    "Google. Material Design 3 Guidelines. https://m3.material.io",
    "Android Developers. Guide to App Architecture & WorkManager. https://developer.android.com",
    "Supabase. Documentation: PostgreSQL, Auth, and Row Level Security. https://supabase.com/docs",
    "Program Studi Informatika. Pedoman Pelaksanaan Capstone Project. Universitas Siliwangi, 2026.",
    "Program Studi Informatika. Rencana Pembelajaran Semester (RPS) Capstone Project. Universitas Siliwangi, 2026.",
]
for i, r in enumerate(refs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-1.0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(f"[{i}]  {r}")
    _set_run(run)

# ---------------------------------------------------------------- save
os.makedirs("docs/proposal", exist_ok=True)
doc.save(OUT)
print("Saved:", OUT, os.path.getsize(OUT), "bytes")
